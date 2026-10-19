"""
Import the Google Apps Script Sudarshan data into PostgreSQL (schema app). Repeatable, and run every 15 minutes
during the trial (scripts/refresh-from-google.ps1): imported bookings, direct calls and triage are updated in place
by their Google ID (their Sudarshan ids never change) and removed when gone from the sheet; the clinic calendar is
replaced (the Google app is its master until cut-over, D-025). Settings, rules, lists and plans of action are only
filled when missing: Sudarshan's Settings page is their master. The final cut-over is one more run.
The old AuditLog is NOT imported.

  python scripts/import_sheets.py                 # live Google Sheet (service account C:\\scripts\\gsheet_creds.json)
  python scripts/import_sheets.py --xlsx FILE     # from an .xlsx export instead
  python scripts/import_sheets.py --dry-run       # read and map everything, then roll back

Never imported: passwords (users must set a new one), API keys, the AuditLog, Yarago copies
(patientMaster, followupMaster, emrBilling ...), planning/optimizer scratch tabs.
"""
import argparse
import datetime as dt
import json
import os
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import psycopg
from dotenv import dotenv_values

ROOT = Path(__file__).resolve().parent.parent
CFG = dotenv_values(ROOT / ".env")
SHEET_ID = CFG.get("GSHEET_ID") or "1QjHA-ovT0Z-SoXtvE7ShbBFbNnknKrnlQwct8zalXXE"
CREDS = CFG.get("GSHEET_CREDS") or r"C:\scripts\gsheet_creds.json"

SHEETS = ["Login", "DoctorAvailability", "DoctorLeaveDays", "DoctorAvailabilityExceptions", "Holidays",
          "planningB", "PlanOfActionMaster", "generalSettings", "RuleSettings", "SystemSettings",
          "appointment", "cancel", "old_cancel", "Archive", "appointment_archive", "RemarksHistory",
          "directCall", "urgencyIdentification", "followup", "followupArchive", "CommentsHistory",
          "Broadcast", "EmailAudit", "DailyKPIReport"]

warnings = Counter()


# ------------------------------------------------------------------ reading
def read_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for name in SHEETS:
        if name in wb.sheetnames:
            out[name] = [list(r) for r in wb[name].iter_rows(values_only=True)]
    return out


def read_gsheet():
    import gspread
    # read-only token: the importer never writes to the Google sheet, whatever the account may share
    gc = gspread.service_account(filename=CREDS, scopes=["https://www.googleapis.com/auth/spreadsheets.readonly",
                                                          "https://www.googleapis.com/auth/drive.readonly"])
    sh = gc.open_by_key(SHEET_ID)
    have = {ws.title: ws for ws in sh.worksheets()}
    out = {}
    for name in SHEETS:
        if name in have:
            try:
                out[name] = have[name].get_values(value_render_option="UNFORMATTED_VALUE",
                                                  date_time_render_option="SERIAL_NUMBER")
            except TypeError:
                out[name] = have[name].get_all_values()
    return out


def records(sheet_rows):
    """rows -> list of dicts keyed by header (blank headers dropped), skipping empty rows."""
    if not sheet_rows:
        return []
    hdr = [("" if h is None else str(h).strip()) for h in sheet_rows[0]]
    out = []
    for n, r in enumerate(sheet_rows[1:], start=2):
        if not any(v not in (None, "") for v in r):
            continue
        d = {"_row": n}
        for i, h in enumerate(hdr):
            if h and i < len(r) and h not in d:
                d[h] = r[i]
        out.append(d)
    return out


# ------------------------------------------------------------------ value cleaning
EPOCH = dt.datetime(1899, 12, 30)


def txt(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"\s+", " ", str(v)).strip()
    return None if s.lower() in ("", "null", "none", "undefined", "nan", "-") else s


def to_ts(v):
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime.combine(v, dt.time())
    if isinstance(v, (int, float)) and 20000 < v < 80000:          # Sheets serial date
        return EPOCH + dt.timedelta(days=float(v))
    s = str(v).strip()
    m = re.match(r"^[A-Z][a-z]{2} ([A-Z][a-z]{2} \d{1,2} \d{4})(?: (\d{2}:\d{2}:\d{2}))?", s)   # JS Date.toString()
    if m:
        try:
            return dt.datetime.strptime(m.group(1) + " " + (m.group(2) or "00:00:00"), "%b %d %Y %H:%M:%S")
        except ValueError:
            pass
    s = s.replace("T", " ").split(".")[0].replace("Z", "")
    for f in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
              "%d/%m/%Y", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d-%m-%Y", "%d %b %Y", "%d-%b-%Y"):
        try:
            return dt.datetime.strptime(s, f)
        except ValueError:
            continue
    warnings["unparsed date/time value"] += 1
    return None


def to_date(v):
    t = to_ts(v)
    return t.date() if t else None


def to_time(v):
    if v in (None, ""):
        return None
    if isinstance(v, dt.time):
        return v.replace(microsecond=0)
    if isinstance(v, dt.datetime):
        return v.time().replace(microsecond=0)
    if isinstance(v, (int, float)) and 0 <= v < 1:                  # Sheets time fraction
        secs = round(v * 86400)
        return dt.time(secs // 3600 % 24, secs // 60 % 60)
    s = str(v).strip().upper()
    for f in ("%H:%M:%S", "%H:%M", "%I:%M %p", "%I:%M%p", "%H.%M"):
        try:
            return dt.datetime.strptime(s, f).time()
        except ValueError:
            continue
    t = to_ts(v)
    if t:
        return t.time()
    warnings["unparsed time value"] += 1
    return None


def to_int(v):
    s = txt(v)
    if s is None:
        return None
    m = re.match(r"^-?\d+(\.\d+)?", s)
    return int(float(m.group(0))) if m else None


def to_bool(v):
    s = (txt(v) or "").lower()
    return s in ("true", "yes", "y", "1", "fixed")


def norm_name(s):
    return re.sub(r"[^a-z]", "", re.sub(r"^\s*dr\.?\s*", "", (s or "").lower()))


ROLE_MAP = {"director": "admin", "admin": "admin", "administrator": "admin", "doctor": "doctor",
            "manager": "manager", "optometrist": "optometrist", "optometry": "optometrist",
            "counsellor": "counsellor", "counselor": "counsellor"}


# ------------------------------------------------------------------ import
APPT_COLS = ["appointment_date", "slot_time", "doctor_id", "mrd", "patient_name", "phone", "age_years", "gender",
             "patient_type", "appointment_type", "plan_of_action", "urgency", "family_group", "fixed_slot", "consult_start",
             "tags", "remarks", "status", "source_ref", "confirmed_4day_at", "confirmed_4day_by", "confirmed_1day_at",
             "confirmed_1day_by", "cancelled_at", "cancelled_by", "booked_by", "booked_at", "legacy_id", "callback_on",
             "cancel_list_done_at", "cancel_list_note"]


def run(data, conn, dry_run=False):
    cur = conn.cursor()
    cur.execute("SET LOCAL app.skip_audit = 'on'")
    cur.execute("SET LOCAL app.source = 'import'")
    counts = Counter()
    R = {k: records(v) for k, v in data.items()}

    # ---- users (password never imported; upsert keeps passwords set later)
    for u in R.get("Login", []):
        username = txt(u.get("username"))
        if not username:
            continue
        role = ROLE_MAP.get((txt(u.get("role")) or "").lower(), "staff")
        cur.execute("""INSERT INTO app.users (username, full_name, role, email, legacy_username)
                       VALUES (%s, %s, %s, %s, %s)
                       ON CONFLICT (username) DO UPDATE SET full_name = EXCLUDED.full_name, role = EXCLUDED.role,
                         email = coalesce(EXCLUDED.email, app.users.email), legacy_username = EXCLUDED.legacy_username""",
                    (username, txt(u.get("Name")) or username, role, txt(u.get("Email")), username))
        counts["users"] += 1
    cur.execute("SELECT id, username::text, full_name FROM app.users")
    user_by = {}
    for uid, un, fn in cur.fetchall():
        user_by[un.lower()] = uid
        user_by[(fn or "").lower()] = uid

    def user_id(v):
        s = (txt(v) or "").lower()
        return user_by.get(s) if s else None

    # ---- doctors
    names = set()
    for sheet in ("DoctorAvailability", "DoctorLeaveDays", "DoctorAvailabilityExceptions", "appointment", "cancel",
                  "old_cancel", "Archive", "appointment_archive"):
        for r in R.get(sheet, []):
            n = txt(r.get("Doctor"))
            if n:
                names.add(n)
    cur.execute("SELECT adoctor_id, tdoctorname FROM yarago.tbldoctor")
    yarago_docs = {norm_name(n): i for i, n in cur.fetchall() if n}
    cur.execute("SELECT name FROM app.doctors")
    canonical = {norm_name(n): n for (n,) in cur.fetchall()}       # one doctor per normalised name
    for n in sorted(names, key=lambda x: (len(x), x)):
        canonical.setdefault(norm_name(n), n)
    for n in sorted(set(canonical.values())):
        cur.execute("""INSERT INTO app.doctors (name, yarago_doctor_id) VALUES (%s, %s)
                       ON CONFLICT (name) DO UPDATE SET yarago_doctor_id = coalesce(app.doctors.yarago_doctor_id, EXCLUDED.yarago_doctor_id)""",
                    (n, yarago_docs.get(norm_name(n))))
    cur.execute("SELECT id, name FROM app.doctors")
    doctor_by = {}
    for i, n in cur.fetchall():
        doctor_by[n.lower()] = i
        doctor_by.setdefault(norm_name(n), i)

    def doctor_id(v):
        s = txt(v)
        return (doctor_by.get(s.lower()) or doctor_by.get(norm_name(s))) if s else None
    counts["doctors"] = len(set(canonical.values()))

    # ---- calendar (replaced wholesale until cut-over)
    cur.execute("DELETE FROM app.doctor_schedule")
    days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
    for r in R.get("DoctorAvailability", []):
        d = doctor_id(r.get("Doctor"))
        for wd, day in enumerate(days, start=1):
            cell = txt(r.get(day))
            if not d or not cell:
                continue
            if cell.lower() in ("off", "closed", "holiday", "leave", "no", "na", "n/a", "not available"):
                continue
            spans = re.findall(r"(\d{1,2}[:.]\d{2})\s*(?:-|–|to)\s*(\d{1,2}[:.]\d{2})", cell, re.I)
            if not spans:
                warnings[f"DoctorAvailability cell not understood: '{cell[:30]}'"] += 1
            for a, b in spans:
                ta, tb = to_time(a.replace(".", ":")), to_time(b.replace(".", ":"))
                if ta and tb and tb > ta:
                    cur.execute("INSERT INTO app.doctor_schedule (doctor_id, weekday, start_time, end_time) VALUES (%s,%s,%s,%s)",
                                (d, wd, ta, tb))
                    counts["doctor_schedule"] += 1
    cur.execute("DELETE FROM app.doctor_exceptions")
    for r in R.get("DoctorLeaveDays", []):
        d, day = doctor_id(r.get("Doctor")), to_date(r.get("Date"))
        if d and day:
            cur.execute("INSERT INTO app.doctor_exceptions (doctor_id, on_date, kind, reason) VALUES (%s,%s,'leave',%s)",
                        (d, day, " - ".join(x for x in (txt(r.get("Reason")), txt(r.get("Notes"))) if x) or None))
            counts["doctor_exceptions"] += 1
    for r in R.get("DoctorAvailabilityExceptions", []):
        d, day = doctor_id(r.get("Doctor")), to_date(r.get("Date"))
        if d and day:
            a, b = to_time(r.get("From Time")), to_time(r.get("To Time"))
            cur.execute("INSERT INTO app.doctor_exceptions (doctor_id, on_date, kind, start_time, end_time, reason) VALUES (%s,%s,%s,%s,%s,%s)",
                        (d, day, "partial" if a and b else "leave", a, b, txt(r.get("Notes"))))
            counts["doctor_exceptions"] += 1
    for r in R.get("Holidays", []):
        day = to_date(r.get("Date"))
        if day:
            cur.execute("""INSERT INTO app.holidays (on_date, reason, remarks) VALUES (%s,%s,%s)
                           ON CONFLICT (on_date) DO UPDATE SET reason = EXCLUDED.reason, remarks = EXCLUDED.remarks""",
                        (day, txt(r.get("Holiday Reason")) or "Holiday", txt(r.get("Remarks"))))
            counts["holidays"] += 1
    for r in R.get("planningB", []):
        day, mx = to_date(r.get("Date")), to_int(r.get("maxSlots"))
        if day and mx is not None:
            cur.execute("""INSERT INTO app.day_capacity (on_date, max_slots, notes) VALUES (%s,%s,%s)
                           ON CONFLICT (on_date) DO UPDATE SET max_slots = EXCLUDED.max_slots""",
                        (day, mx, txt(r.get("maxSlot edit history"))))
            counts["day_capacity"] += 1
    for r in R.get("PlanOfActionMaster", []):
        n = txt(r.get("Plan of action"))
        if n:
            cur.execute("""INSERT INTO app.plan_of_action (name, eye, prep_min, consult_min, priority) VALUES (%s,%s,%s,%s,%s)
                           ON CONFLICT (name) DO NOTHING""",              # Sudarshan's Settings page owns plans
                        (n, txt(r.get("Eye")), to_int(r.get("PrepMin")), to_int(r.get("ConsultMin")), to_int(r.get("Priority"))))
            counts["plan_of_action"] += 1

    # ---- settings: only filled when missing; Sudarshan's Settings page owns them (D-025)
    def put_setting(key, value, desc=None):
        cur.execute("""INSERT INTO app.settings (key, value, description) VALUES (%s, %s::jsonb, %s)
                       ON CONFLICT (key) DO NOTHING""",
                    (key, json.dumps(value, default=str), desc))
        counts["settings"] += 1
    gs = data.get("generalSettings") or []
    if gs:
        hdr = gs[0]
        for i, h in enumerate(hdr):
            h = txt(h)
            if h:
                vals = [txt(r[i]) for r in gs[1:] if i < len(r) and txt(r[i])]
                put_setting("general." + h, vals, "Imported from generalSettings")
    for r in R.get("RuleSettings", []):
        k = txt(r.get("ruleKey"))
        if k:
            put_setting("rule." + k, {"name": txt(r.get("ruleName")), "value": txt(r.get("value"))}, txt(r.get("description")))
    for r in R.get("SystemSettings", [])[:1]:
        for k, v in r.items():
            if k != "_row" and txt(v):
                put_setting("system." + k, txt(v), "Imported from SystemSettings")

    # ---- appointments: one row per AppointmentID across all appointment sheets, updated in place by that id.
    # The import's own history events (actor_id empty) are rewritten; anything a Sudarshan user added is kept.
    cur.execute("""DELETE FROM app.appointment_events WHERE actor_id IS NULL
                     AND appointment_id IN (SELECT id FROM app.appointments WHERE legacy_id IS NOT NULL)""")
    cur.execute("SELECT legacy_id, id FROM app.appointments WHERE legacy_id IS NOT NULL")
    appt_ids = dict(cur.fetchall())
    seen, events = set(), []
    order = [("appointment", "active"), ("cancel", "cancelled"), ("old_cancel", "old_cancel"),
             ("Archive", "archived_cancel"), ("appointment_archive", "past")]
    for sheet, kind in order:
        for r in R.get(sheet, []):
            lid = txt(r.get("AppointmentID")) or f"noid:{sheet}:{r['_row']}"
            if lid in seen:
                warnings[f"duplicate AppointmentID skipped ({sheet})"] += 1
                continue
            day = to_date(r.get("AppointmentDate"))
            name = txt(r.get("PatientName"))
            if not day or not name:
                warnings[f"{sheet}: row without date or name skipped"] += 1
                continue
            seen.add(lid)
            tags = [t.strip() for t in re.split(r"[,;]", txt(r.get("Tag")) or "") if t.strip()]
            slot_status = (txt(r.get("SlotStatus")) or "").lower()
            slot = to_time(r.get("TimeSlot"))
            if slot is not None and slot.hour == 0 and slot.minute == 0:
                slot = None                     # the Google app's "no time" (an empty cell read as midnight): waitlist
            if kind == "active":
                status = "waitlist" if ("wait" in slot_status or slot is None) else (
                    "confirmed" if r.get("OneDayConfirmDate") or r.get("FourDayConfirmDate") else "booked")
            elif kind == "past":
                status = "completed"            # refined below against Yarago arrivals
            else:
                status = "cancelled"
            if kind == "archived_cancel":
                tags.append("archived from cancel list")
            call_back = txt(r.get("call back"))
            if call_back:
                tags.append("call back: " + call_back)
            src_ref = txt(r.get("aFollowID")) or txt(r.get("aFollowID / directCallID"))
            vals = dict(zip(APPT_COLS, (day, slot, doctor_id(r.get("Doctor")), (txt(r.get("MRDNo")) or "").upper() or None, name,
               txt(r.get("Phone")), to_int(r.get("Age")), txt(r.get("Gender")),
               txt(r.get("PatientType")) or txt(r.get("Patient Type")), txt(r.get("AppointmentType")), txt(r.get("PlanOfAction")),
               txt(r.get("Urgency Level")) or txt(r.get("Urgency level (waiting List)")), txt(r.get("Family Identifier")),
               to_bool(r.get("Fixed Slot") or r.get("fixed slot")), to_time(r.get("Consult Start")), tags, txt(r.get("Remarks")),
               status, src_ref, to_ts(r.get("FourDayConfirmDate")), user_id(r.get("FourDayConfirmBy")),
               to_ts(r.get("OneDayConfirmDate")), user_id(r.get("OneDayConfirmBy")),
               to_ts(r.get("Deleted On")) if status == "cancelled" else None,
               user_id(r.get("Deleted by")) if status == "cancelled" else None,
               user_id(r.get("BookedBy")), to_ts(r.get("BookingDate")) or dt.datetime.now(), lid,
               # cancel list (0007): call-back date, and rows archived from the cancel list are off it
               to_date(r.get("Call back date")) if status == "cancelled" else None,
               (to_ts(r.get("Archive date")) or to_ts(r.get("Archived on")) or to_ts(r.get("Deleted On")) or dt.datetime.now())
               if kind in ("archived_cancel", "old_cancel") else None,
               {"archived_cancel": "Archived from the cancel list in the Google app",
                "old_cancel": "Old cancellation (old_cancel sheet in the Google app)"}.get(kind))))
            # update a booking already copied (same Sudarshan id, no id number used up); insert only new ones.
            # When the sheet has no booking/archive time, keep the one stored rather than "now" on every run.
            if lid in appt_ids:
                sets = [f"{c} = %({c})s" for c in APPT_COLS if c not in ("booked_at", "cancel_list_done_at", "legacy_id")]
                sets.append("booked_at = " + ("%(booked_at)s" if to_ts(r.get("BookingDate")) else "coalesce(booked_at, %(booked_at)s)"))
                sets.append("cancel_list_done_at = CASE WHEN %(cancel_list_done_at)s::timestamptz IS NULL THEN NULL "
                            "ELSE coalesce(cancel_list_done_at, %(cancel_list_done_at)s) END")
                cur.execute(f"""UPDATE app.appointments SET
                                  patient_id = CASE WHEN mrd IS DISTINCT FROM %(mrd)s THEN NULL ELSE patient_id END,
                                  yarago_visit_id = CASE WHEN mrd IS DISTINCT FROM %(mrd)s
                                                           OR appointment_date IS DISTINCT FROM %(appointment_date)s
                                                         THEN NULL ELSE yarago_visit_id END,
                                  {", ".join(sets)}
                                WHERE id = %(id)s RETURNING id""", {**vals, "id": appt_ids[lid]})
            else:
                cur.execute(f"""INSERT INTO app.appointments (source, {", ".join(APPT_COLS)})
                                VALUES ('import', {", ".join(f"%({c})s" for c in APPT_COLS)}) RETURNING id""", vals)
            new_id = cur.fetchone()[0]
            counts[f"appointments ({sheet})"] += 1
            history = {k: txt(r.get(k)) for k in ("RescheduleHistory", "AppointmentJourney", "Call back date",
                                                  "Archived on", "Archived by", "Archive date", "Archived By") if txt(r.get(k))}
            history["sheet"] = sheet
            note = "; ".join(x for x in (
                f"booked by {txt(r.get('BookedBy'))}" if txt(r.get("BookedBy")) else None,
                f"rescheduled by {txt(r.get('RescheduledBy'))} on {txt(r.get('RescheduledDate'))}" if txt(r.get("RescheduledBy")) else None,
                f"cancelled by {txt(r.get('Deleted by'))}" if status == "cancelled" and txt(r.get("Deleted by")) else None) if x)
            events.append((new_id, to_ts(r.get("BookingDate")) or dt.datetime.now(), "imported", json.dumps(history, default=str),
                           f"Imported from Google Sheets ({sheet})" + (f": {note}" if note else "")))

    # bookings no longer in the sheet (deleted in the Google app) go here too
    cur.execute("SELECT id FROM app.appointments WHERE legacy_id IS NOT NULL AND NOT (legacy_id = ANY(%s))", (list(seen),))
    gone = [r[0] for r in cur.fetchall()]
    if gone:
        cur.execute("UPDATE app.direct_calls SET appointment_id = NULL WHERE appointment_id = ANY(%s)", (gone,))
        cur.execute("DELETE FROM app.appointment_events WHERE appointment_id = ANY(%s)", (gone,))
        cur.execute("DELETE FROM app.appointments WHERE id = ANY(%s)", (gone,))
    counts["bookings removed (no longer in the sheet)"] = len(gone)

    cur.execute("SELECT id, legacy_id FROM app.appointments WHERE legacy_id IS NOT NULL")
    appt_by = dict((l, i) for i, l in cur.fetchall())
    for r in R.get("RemarksHistory", []):
        a = appt_by.get(txt(r.get("AppointmentID")) or "")
        if a and txt(r.get("Remark")):
            events.append((a, to_ts(r.get("AddedAt")) or dt.datetime.now(), "remark", None, txt(r.get("Remark")),
                           txt(r.get("AddedBy"))))
        elif txt(r.get("Remark")):
            warnings["remark for an appointment not found"] += 1
    with cur.copy("COPY app.appointment_events (appointment_id, at, event_type, to_value, note, actor_name) FROM STDIN") as cp:
        for e in events:
            cp.write_row(e if len(e) == 6 else e + (None,))
    counts["appointment_events"] = len(events)

    # ---- direct calls, triage
    seen_calls = []
    for r in R.get("directCall", []):
        st = (txt(r.get("Status")) or "").lower()
        status = "booked" if ("book" in st or "appoint" in st) else "cancelled" if "cancel" in st else "archived" if "archiv" in st else "open"
        cur.execute("""INSERT INTO app.direct_calls (mrd, patient_name, phone, age_years, gender, patient_type, call_type, case_type,
                         urgency, remarks, status, created_by, created_at, legacy_id)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                       ON CONFLICT (legacy_id) DO UPDATE SET mrd = EXCLUDED.mrd, patient_name = EXCLUDED.patient_name,
                         phone = EXCLUDED.phone, age_years = EXCLUDED.age_years, gender = EXCLUDED.gender, patient_type = EXCLUDED.patient_type,
                         call_type = EXCLUDED.call_type, case_type = EXCLUDED.case_type, urgency = EXCLUDED.urgency,
                         remarks = EXCLUDED.remarks, status = EXCLUDED.status, created_by = EXCLUDED.created_by, created_at = EXCLUDED.created_at""",
                    ((txt(r.get("MRDNo")) or "").upper() or None, txt(r.get("PatientName")), txt(r.get("Phone")), to_int(r.get("Age")),
                     txt(r.get("Gender")), txt(r.get("PatientType")), txt(r.get("CallType")), txt(r.get("CaseType")),
                     txt(r.get("Urgency")), txt(r.get("Remarks")), status, user_id(r.get("CreatedBy")),
                     to_ts(r.get("CreatedOn")) or dt.datetime.now(), txt(r.get("DirectCallID")) or f"noid:directCall:{r['_row']}"))
        seen_calls.append(txt(r.get("DirectCallID")) or f"noid:directCall:{r['_row']}")
        counts["direct_calls"] += 1
    cur.execute("DELETE FROM app.direct_calls WHERE legacy_id IS NOT NULL AND NOT (legacy_id = ANY(%s))", (seen_calls,))
    counts["direct calls removed (no longer in the sheet)"] = cur.rowcount
    seen_triage = []
    qmap = {"Q1_Onset": "onset", "Q2_Vision": "vision", "Q3_Pain": "pain", "Q4_Floaters": "floaters",
            "Q5_Redness": "redness", "Q6_Injury": "injury", "Q7_Symptoms": "symptoms"}
    for r in R.get("urgencyIdentification", []):
        answers = {v: txt(r.get(k)) for k, v in qmap.items() if txt(r.get(k)) is not None}
        cur.execute("""INSERT INTO app.triage (mrd, patient_name, phone, age_years, gender, patient_type, assessed_at, answers,
                         max_score, category, decision_status, decision_remarks, decision_by, decision_at, created_by, created_at, legacy_id)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s::jsonb,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                       ON CONFLICT (legacy_id) DO UPDATE SET mrd = EXCLUDED.mrd, patient_name = EXCLUDED.patient_name, phone = EXCLUDED.phone,
                         age_years = EXCLUDED.age_years, gender = EXCLUDED.gender, patient_type = EXCLUDED.patient_type,
                         assessed_at = EXCLUDED.assessed_at, answers = EXCLUDED.answers, max_score = EXCLUDED.max_score,
                         category = EXCLUDED.category, decision_status = EXCLUDED.decision_status, decision_remarks = EXCLUDED.decision_remarks,
                         decision_by = EXCLUDED.decision_by, decision_at = EXCLUDED.decision_at, created_by = EXCLUDED.created_by""",
                    ((txt(r.get("MRDNo")) or "").upper() or None, txt(r.get("PatientName")), txt(r.get("Phone")), to_int(r.get("Age")),
                     txt(r.get("Gender")), txt(r.get("PatientType")), to_ts(r.get("AssessmentDate")) or to_ts(r.get("CreatedOn")) or dt.datetime.now(),
                     json.dumps(answers), to_int(r.get("MaxScore")), txt(r.get("Category")), txt(r.get("DecisionStatus")),
                     txt(r.get("DecisionRemarks")), user_id(r.get("DecisionBy")), to_ts(r.get("DecisionDate")),
                     user_id(r.get("CreatedBy")), to_ts(r.get("CreatedOn")) or dt.datetime.now(),
                     txt(r.get("TriageID")) or f"noid:triage:{r['_row']}"))
        seen_triage.append(txt(r.get("TriageID")) or f"noid:triage:{r['_row']}")
        counts["triage"] += 1
    cur.execute("DELETE FROM app.triage WHERE legacy_id IS NOT NULL AND NOT (legacy_id = ANY(%s))", (seen_triage,))
    counts["triage removed (no longer in the sheet)"] = cur.rowcount

    # ---- follow-up work history
    cur.execute("DELETE FROM app.followup_actions WHERE legacy_ref IS NOT NULL")
    fa = []
    for r in R.get("CommentsHistory", []):
        if txt(r.get("Comment")):
            fa.append((to_int(r.get("aFollowID")), None, "note", txt(r.get("Comment")), txt(r.get("AddedBy")),
                       to_ts(r.get("AddedAt")) or dt.datetime.now(), f"CommentsHistory:{r['_row']}"))
    for sheet in ("followup", "followupArchive"):
        for r in R.get(sheet, []):
            parts = [f"{k}: {txt(r.get(k))}" for k in ("Status", "Tags", "Source", "ArchiveReason") if txt(r.get(k))]
            if parts:
                fa.append((to_int(r.get("aFollowID")), to_int(r.get("npatient_id")), "note",
                           ("Archived. " if sheet == "followupArchive" else "") + "; ".join(parts), None,
                           to_ts(r.get("ArchivedOn")) or to_ts(r.get("updated_date")) or dt.datetime.now(), f"{sheet}:{r['_row']}"))
    with cur.copy("COPY app.followup_actions (yarago_followup_id, patient_id, action, note, actor_name, at, legacy_ref) FROM STDIN") as cp:
        for row in fa:
            cp.write_row(row)
    counts["followup_actions"] = len(fa)

    # ---- messages (broadcast summaries, email log), daily KPIs
    cur.execute("DELETE FROM app.messages WHERE legacy_ref IS NOT NULL")
    for r in R.get("Broadcast", []):
        cur.execute("""INSERT INTO app.messages (channel, template, recipient, payload, status, created_at, sent_at, legacy_ref)
                       VALUES ('whatsapp', %s, 'broadcast', %s::jsonb, 'sent', %s, %s, %s)""",
                    (txt(r.get("CampaignName")), json.dumps({"appointment_date": str(to_date(r.get("AppointmentDate"))),
                     "sent": to_int(r.get("SentCount")), "skipped": to_int(r.get("SkippedCount")),
                     "failed": to_int(r.get("FailedCount")), "sent_by": txt(r.get("SentBy"))}),
                     to_ts(r.get("Timestamp")) or dt.datetime.now(), to_ts(r.get("Timestamp")), f"Broadcast:{r['_row']}"))
        counts["messages"] += 1
    for r in R.get("EmailAudit", []):
        if txt(r.get("Recipient")):
            ok = "fail" not in (txt(r.get("Status")) or "").lower()
            cur.execute("""INSERT INTO app.messages (channel, template, recipient, payload, status, created_at, sent_at, legacy_ref)
                           VALUES ('email', 'daily report', %s, %s::jsonb, %s, %s, %s, %s)""",
                        (txt(r.get("Recipient")), json.dumps({"for_date": str(to_date(r.get("Date"))), "sent_by": txt(r.get("Sent By"))}),
                         "sent" if ok else "failed", to_ts(r.get("Timestamp")) or dt.datetime.now(), to_ts(r.get("Timestamp")),
                         f"EmailAudit:{r['_row']}"))
            counts["messages"] += 1
    for r in R.get("DailyKPIReport", []):
        day = to_date(r.get("Date"))
        if day:
            metrics = {k: (to_int(v) if to_int(v) is not None and re.match(r"^-?\d+(\.0)?$", str(v).strip()) else txt(v))
                       for k, v in r.items() if k not in ("_row", "Date", "Timestamp")}
            metrics["source"] = "legacy DailyKPIReport"
            cur.execute("""INSERT INTO app.kpi_daily (on_date, metrics, computed_at) VALUES (%s, %s::jsonb, %s)
                           ON CONFLICT (on_date) DO UPDATE SET metrics = EXCLUDED.metrics, computed_at = EXCLUDED.computed_at""",
                        (day, json.dumps(metrics, default=str), to_ts(r.get("Timestamp")) or dt.datetime.now()))
            counts["kpi_daily"] += 1

    # ---- link to Yarago patients by MRD, and tell completed from no-show for past bookings
    for table in ("appointments", "direct_calls", "triage"):
        cur.execute(f"""UPDATE app.{table} t SET patient_id = p.apatient_id
                        FROM yarago.tblpatientregistration p
                        WHERE t.patient_id IS NULL AND t.mrd IS NOT NULL AND upper(btrim(p.apatientcode)) = t.mrd""")
        counts[f"{table} linked to Yarago patient"] = cur.rowcount
    cur.execute("""UPDATE app.appointments a SET status = 'no_show'
                   WHERE a.legacy_id IS NOT NULL AND a.status = 'completed' AND a.patient_id IS NOT NULL
                     AND a.appointment_date < current_date
                     AND NOT EXISTS (SELECT 1 FROM yarago.tblappointments y
                                     WHERE y.npatient_id = a.patient_id AND y.tappoint_date_time::date = a.appointment_date)""")
    counts["past bookings with no Yarago visit that day -> no_show"] = cur.rowcount
    cur.execute("""UPDATE app.appointments a SET yarago_visit_id = y.aappoint_id
                   FROM yarago.tblappointments y
                   WHERE a.legacy_id IS NOT NULL AND a.yarago_visit_id IS NULL AND y.npatient_id = a.patient_id
                     AND y.tappoint_date_time::date = a.appointment_date""")
    counts["bookings matched to a Yarago visit"] = cur.rowcount

    cur.execute("INSERT INTO app.import_runs (source, finished_at, status, counts, message) VALUES (%s, now(), %s, %s::jsonb, %s)",
                ("google_sheets", "dry-run" if dry_run else "ok", json.dumps(counts), json.dumps(dict(warnings))[:4000]))
    return counts


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", help="read an .xlsx export instead of the live sheet")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--scheduled", action="store_true",
                    help="trial refresh: does nothing once the setting trial.google_master is false (after cut-over)")
    a = ap.parse_args()
    if a.scheduled:
        with psycopg.connect(host=CFG["PG_HOST"], port=int(CFG["PG_PORT"]), dbname=CFG["PG_DB"],
                             user=CFG["PG_APP_USER"], password=CFG["PG_APP_PASSWORD"]) as conn:
            row = conn.execute("SELECT value FROM app.settings WHERE key = 'trial.google_master'").fetchone()
        if not row or row[0] is not True:
            print("trial.google_master is not true: Sudarshan is the master now, nothing imported")
            return 0
    data = read_xlsx(a.xlsx) if a.xlsx else read_gsheet()
    print("sheets read:", ", ".join(f"{k} ({max(len(v) - 1, 0)})" for k, v in data.items()))
    missing = [s for s in SHEETS if s not in data]
    if missing:
        print("not found (skipped):", ", ".join(missing))
    with psycopg.connect(host=CFG["PG_HOST"], port=int(CFG["PG_PORT"]), dbname=CFG["PG_DB"],
                         user=CFG["PG_APP_USER"], password=CFG["PG_APP_PASSWORD"]) as conn:
        counts = run(data, conn, a.dry_run)
        if a.dry_run:
            conn.rollback()
        else:
            conn.commit()
    print("\nimported" + (" (DRY RUN, rolled back)" if a.dry_run else "") + ":")
    for k, v in counts.items():
        print(f"  {k:55} {v:>8,}")
    if warnings:
        print("\nwarnings:")
        for k, v in warnings.most_common():
            print(f"  {v:>6,} x {k}")


if __name__ == "__main__":
    sys.exit(main() or 0)
